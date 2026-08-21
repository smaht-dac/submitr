"""Submitr integration tests for dcicutils ProtectedDonor transformation."""

from pathlib import Path
from unittest.mock import Mock, patch
import json
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import pytest

from dcicutils.common import APP_FOURFRONT
from dcicutils.submitr.custom_excel import CustomExcel
from dcicutils.submitr import donor_transformer as _donor_transformer_module

if not hasattr(_donor_transformer_module, "DonorReferenceKind"):
    # The pinned dcicutils dependency (4dn-dcic/utils branch ajs/protected-donor-transform,
    # tracked at https://github.com/4dn-dcic/utils/pull/337) has not yet landed the analyze()/
    # DonorReferenceKind/portal-aware ProtectedDonor API these tests exercise. Skip rather than
    # fail collection until that dependency catches up.
    pytest.skip(
        "dcicutils.submitr.donor_transformer.DonorReferenceKind is not available in the "
        "pinned dcicutils dependency; see 4dn-dcic/utils PR #337.",
        allow_module_level=True,
    )

from dcicutils.submitr.donor_transformer import (
    DonorReferenceKind,
    ProtectedDonorTransformError,
    ProtectedDonorWorkbookTransformer,
)

from ..submission import (
    _analyze_protected_donor_workbook,
    _ensure_protected_donor_transformed_workbook,
    _initiate_server_ingestion_process,
    _validate_locally,
    _prepare_protected_donor_transform,
    _pre_transform_to_temp_json,
    _stage_protected_donor_workbook,
    submit_any_ingestion,
)


PROTECTED_SHEETS = [
    "Demographic",
    "DeathCircumstances",
    "FamilyHistory",
    "MedicalHistory",
    "TissueCollection",
]


def make_workbook(path: Path, *, references=None, donor_rows=None, protected_rows=None,
                  extra_reference_sheet=None):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    donor = workbook.create_sheet("Donor")
    donor.append(["submitted_id", "name", "status"])
    for row in donor_rows or []:
        donor.append(row)
    if protected_rows is not None:
        protected = workbook.create_sheet("ProtectedDonor")
        protected.append(["submitted_id", "name", "status"])
        for row in protected_rows:
            protected.append(row)
    for sheet_name in PROTECTED_SHEETS:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(["donor", "value"])
        for reference in (references or {}).get(sheet_name, []):
            sheet.append([reference, "x"])
    if extra_reference_sheet:
        sheet = workbook.create_sheet(extra_reference_sheet[0])
        sheet.append(["donor"])
        sheet.append([extra_reference_sheet[1]])
    workbook.save(path)
    return path


def add_formula_cached_value(path: Path, *, value=None):
    workbook = openpyxl.load_workbook(path)
    workbook["Demographic"]["A2"] = '=\"D_DONOR_1\"'
    workbook.save(path)
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    namespaces = {"x": namespace}
    with zipfile.ZipFile(path) as source:
        files = {name: source.read(name) for name in source.namelist()}
    for name, contents in list(files.items()):
        if not (name.startswith("xl/worksheets/sheet") and name.endswith(".xml")):
            continue
        sheet = ET.fromstring(contents)
        for cell in sheet.findall(".//x:c", namespaces):
            if cell.find("x:f", namespaces) is None:
                continue
            if value is not None:
                cell.set("t", "str")
                cached = cell.find("x:v", namespaces)
                cached.text = value
        files[name] = ET.tostring(sheet, encoding="utf-8", xml_declaration=True)
    with zipfile.ZipFile(path, "w") as destination:
        for name, contents in files.items():
            destination.writestr(name, contents)


def test_transform_only_referenced_donors_and_only_five_protected_sheets(tmp_path):
    path = make_workbook(
        tmp_path / "input.xlsx",
        donor_rows=[["D_DONOR_1", "referenced", "active"],
                    ["D_DONOR_2", "untouched", "active"]],
        references={"Demographic": ["D_DONOR_1"]},
        extra_reference_sheet=("Other", "D_DONOR_2"),
    )
    workbook = openpyxl.load_workbook(path)
    changed = ProtectedDonorWorkbookTransformer().transform(workbook)

    assert changed is True
    donor = workbook["Donor"]
    headers = [cell.value for cell in donor[1]]
    assert donor.cell(2, headers.index("protected_donor") + 1).value == "D_PROTECTED-DONOR_1"
    assert all(cell.value != "D_PROTECTED-DONOR_2" for cell in donor[3])
    assert workbook["Demographic"]["A2"].value == "D_PROTECTED-DONOR_1"
    assert workbook["Other"]["A2"].value == "D_DONOR_2"
    assert workbook["ProtectedDonor"].max_row == 2


def test_mixed_references_are_analyzed_and_transformed_incrementally(tmp_path):
    path = make_workbook(
        tmp_path / "mixed.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"], ["D_DONOR_2", "two", "active"]],
        protected_rows=[["D_PROTECTED-DONOR_2", "two", "in review"]],
        references={"Demographic": ["D_DONOR_1", "D_PROTECTED-DONOR_2"]},
    )
    workbook = openpyxl.load_workbook(path)
    analysis = ProtectedDonorWorkbookTransformer().analyze(workbook)
    assert analysis.references["D_DONOR_1"] == DonorReferenceKind.PLAIN_DONOR
    assert analysis.references["D_PROTECTED-DONOR_2"] == DonorReferenceKind.PROTECTED_DONOR
    assert analysis.needs_transformation

    ProtectedDonorWorkbookTransformer().transform(workbook)
    assert workbook["Demographic"]["A2"].value == "D_PROTECTED-DONOR_1"
    assert workbook["Demographic"]["A3"].value == "D_PROTECTED-DONOR_2"
    assert workbook["ProtectedDonor"].max_row == 3

    # A second pass is incremental: it does not duplicate existing rows.
    assert ProtectedDonorWorkbookTransformer().transform(workbook) is False
    assert workbook["ProtectedDonor"].max_row == 3


def test_missing_workbook_donor_is_an_error(tmp_path):
    path = tmp_path / "missing.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Demographic"
    workbook.active.append(["donor"])
    workbook.active.append(["D_DONOR_MISSING"])
    workbook.save(path)

    with pytest.raises(ProtectedDonorTransformError, match="absent from the workbook Donor sheet"):
        transformer = ProtectedDonorWorkbookTransformer()
        transformer.transform(openpyxl.load_workbook(path))


def test_protected_donor_can_be_resolved_from_workbook_or_portal(tmp_path):
    path = make_workbook(
        tmp_path / "portal.xlsx",
        references={"Demographic": ["D_PROTECTED-DONOR_PORTAL"]},
        protected_rows=None,
    )
    portal = Mock()
    portal.ref_exists.return_value = True
    analysis = _analyze_protected_donor_workbook(path, portal)
    assert analysis.references["D_PROTECTED-DONOR_PORTAL"] == DonorReferenceKind.PROTECTED_DONOR
    assert not analysis.needs_transformation


def test_submitr_analysis_distinguishes_partial_and_invalid_workbooks(tmp_path):
    partial = make_workbook(
        tmp_path / "partial.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"]],
        protected_rows=[["D_PROTECTED-DONOR_2", "two", "in review"]],
        references={"Demographic": ["D_DONOR_1", "D_PROTECTED-DONOR_2"]},
    )
    invalid = make_workbook(
        tmp_path / "invalid.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"]],
        references={"Demographic": ["not-an-id"]},
    )
    assert _prepare_protected_donor_transform(
        portal=None, ingestion_filename=str(partial), submission_centers=None,
        validation=False, no_query=False, transform_protected_donor=True,
        transformed_workbook_path=None,
    )[0] is True
    invalid_analysis = _analyze_protected_donor_workbook(invalid, None)
    assert invalid_analysis.invalid_references == frozenset({"not-an-id"})
    with pytest.raises(ProtectedDonorTransformError, match="Invalid donor reference"):
        _prepare_protected_donor_transform(
            portal=None, ingestion_filename=str(invalid), submission_centers=None,
            validation=False, no_query=False, transform_protected_donor=True,
            transformed_workbook_path=None,
        )


def test_prepare_auto_transforms_without_prompt_without_printing_output_path(tmp_path):
    path = make_workbook(
        tmp_path / "input.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"]],
        references={"Demographic": ["D_DONOR_1"]},
    )
    with patch("submitr.submission.yes_or_no") as prompt, \
            patch("submitr.submission.PRINT") as output:
        apply, output_path = _prepare_protected_donor_transform(
            portal=None, ingestion_filename=str(path), submission_centers=None,
            validation=False, no_query=False, transform_protected_donor=True,
            transformed_workbook_path=None,
        )
    assert apply is True
    assert output_path.endswith("input.transformed.xlsx")
    prompt.assert_not_called()
    messages = [str(call.args[0]) for call in output.call_args_list]
    assert not any("transformation is occurring" in message for message in messages)
    assert not any("Transformed workbook path:" in message for message in messages)


def test_custom_excel_saves_the_transformed_workbook(tmp_path):
    source = make_workbook(
        tmp_path / "source.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"]],
        references={"Demographic": ["D_DONOR_1"]},
    )
    destination = tmp_path / "source.transformed.xlsx"
    CustomExcel.with_portal(
        None, transform_protected_donor=True,
        transformed_workbook_path=str(destination),
    )(str(source))
    transformed = openpyxl.load_workbook(destination)
    assert transformed["Demographic"]["A2"].value == "D_PROTECTED-DONOR_1"


def test_custom_excel_protects_caller_selected_existing_output(tmp_path):
    source = make_workbook(
        tmp_path / "source.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"]],
        references={"Demographic": ["D_DONOR_1"]},
    )
    destination = tmp_path / "source.transformed.xlsx"
    previous = b"caller-owned output"
    destination.write_bytes(previous)

    with pytest.raises(ValueError, match="output path already exists"):
        CustomExcel.with_portal(
            None, transform_protected_donor=True,
            transformed_workbook_path=str(destination),
        )(str(source))

    assert destination.read_bytes() == previous


def test_local_validation_replaces_existing_owned_staging_file_after_success(tmp_path):
    source = make_workbook(
        tmp_path / "source.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"]],
        references={"Demographic": ["D_DONOR_1"]},
    )
    destination = tmp_path / "source.transformed.xlsx"
    destination.write_bytes(b"pre-created staging file")

    def stage_with_existing_path(path):
        staged_path = _stage_protected_donor_workbook(path)
        assert Path(staged_path).exists()
        return staged_path

    with patch("submitr.submission._validate_data", return_value=True), \
            patch("submitr.submission._stage_protected_donor_workbook", side_effect=stage_with_existing_path), \
            pytest.raises(SystemExit) as exit_info:
        _validate_locally(
            str(source), None, autoadd={}, validation=True,
            validate_local_only=True, exit_immediately_on_errors=True,
            noprogress=True, noanalyze=True, transform_protected_donor=True,
            transformed_workbook_path=str(destination),
        )

    assert exit_info.value.code == 0
    assert openpyxl.load_workbook(destination)["Demographic"]["A2"].value == (
        "D_PROTECTED-DONOR_1"
    )


def test_protected_donor_remote_payload_contains_mapped_data(tmp_path):
    source = make_workbook(
        tmp_path / "source.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"]],
        references={"Demographic": ["D_DONOR_1"]},
    )
    workbook = openpyxl.load_workbook(source)
    quality_metrics = workbook.create_sheet("ExternalQualityMetric")
    quality_metrics.append(["submitted_id", "total_raw_reads_sequenced"])
    quality_metrics.append(["EQM_1", 123])
    workbook.save(source)

    transformed_path = tmp_path / "source.transformed.xlsx"
    structured_data = _ensure_protected_donor_transformed_workbook(
        str(source), None, str(transformed_path)
    )
    upload_path = _pre_transform_to_temp_json(str(source), structured_data)
    try:
        portal = Mock()
        portal.server = "https://example.test"
        created = Mock()
        created.json.return_value = {"@graph": [{"@id": "/ingestion-submissions/1"}]}
        initiated = Mock()
        initiated.json.return_value = {"submission_id": "submission-1"}
        portal.post.side_effect = [created, initiated]
        _initiate_server_ingestion_process(
            portal=portal,
            ingestion_filename=str(transformed_path),
            upload_filename=upload_path,
            consortia=[],
            submission_centers=[],
        )
        uploaded = portal.post.call_args_list[1].kwargs["files"]["datafile"]
        payload = json.loads(uploaded.read().decode("utf-8"))
        uploaded.close()
    finally:
        if upload_path:
            Path(upload_path).unlink(missing_ok=True)

    assert payload["Demographic"][0]["donor"] == "D_PROTECTED-DONOR_1"
    assert payload["ExternalQualityMetric"][0]["qc_values"] == [{
        "derived_from": "total_raw_reads_sequenced",
        "value": 123,
        "key": "Total Raw Reads Sequenced",
        "tooltip": "# of reads (150bp)",
    }]
    assert openpyxl.load_workbook(transformed_path)["Demographic"]["A2"].value == "D_PROTECTED-DONOR_1"


def test_formula_cached_donor_reference_is_used_without_evaluation(tmp_path):
    source = make_workbook(
        tmp_path / "formula.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"]],
        references={"Demographic": ["D_DONOR_1"]},
    )
    add_formula_cached_value(source, value="D_DONOR_1")

    analysis = _analyze_protected_donor_workbook(source, None)
    assert analysis.needs_transformation
    destination = tmp_path / "formula.transformed.xlsx"
    structured_data = _ensure_protected_donor_transformed_workbook(
        str(source), None, str(destination)
    )
    assert structured_data.data["Demographic"][0]["donor"] == "D_PROTECTED-DONOR_1"
    assert (
        openpyxl.load_workbook(destination, data_only=True)["Demographic"]["A2"].value
        == "D_PROTECTED-DONOR_1"
    )


@pytest.mark.parametrize(
    "cached_value, message",
    [(None, "no valid cached value"), ("not-a-donor", "invalid cached value")],
)
def test_formula_donor_reference_requires_a_valid_cached_value(tmp_path, cached_value, message):
    source = make_workbook(
        tmp_path / "formula.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"]],
        references={"Demographic": ["D_DONOR_1"]},
    )
    add_formula_cached_value(source, value=cached_value)

    with pytest.raises(ProtectedDonorTransformError, match=message):
        _analyze_protected_donor_workbook(source, None)


def test_existing_transformed_workbook_is_preserved_until_successful_replacement(tmp_path):
    source = make_workbook(
        tmp_path / "source.xlsx",
        donor_rows=[["D_DONOR_1", "one", "active"]],
        references={"Demographic": ["D_DONOR_1"]},
    )
    destination = tmp_path / "source.transformed.xlsx"
    destination.write_bytes(b"previous transformed workbook")
    previous = destination.read_bytes()

    _prepare_protected_donor_transform(
        portal=None, ingestion_filename=str(source), submission_centers=None,
        validation=False, no_query=False, transform_protected_donor=True,
        transformed_workbook_path=str(destination),
    )
    assert destination.read_bytes() == previous

    _ensure_protected_donor_transformed_workbook(str(source), None, str(destination))
    assert destination.read_bytes() != previous
    assert openpyxl.load_workbook(destination)["Demographic"]["A2"].value == "D_PROTECTED-DONOR_1"


def test_failed_protected_donor_load_preserves_existing_target_and_has_cause(tmp_path):
    source = tmp_path / "malformed.xlsx"
    source.write_bytes(b"not an xlsx workbook")
    destination = tmp_path / "malformed.transformed.xlsx"
    previous = b"previous transformed workbook"
    destination.write_bytes(previous)

    with pytest.raises(ProtectedDonorTransformError, match="Unable to load ProtectedDonor workbook") as error:
        _prepare_protected_donor_transform(
            portal=None, ingestion_filename=str(source), submission_centers=None,
            validation=False, no_query=False, transform_protected_donor=True,
            transformed_workbook_path=str(destination),
        )

    assert error.value.__cause__ is not None
    assert destination.read_bytes() == previous


@pytest.mark.parametrize(
    "mode",
    ["validate_local_only", "validate_remote_only", "validate_local_skip"],
)
def test_malformed_workbook_fails_before_each_validation_mode(tmp_path, mode):
    source = tmp_path / "malformed.xlsx"
    source.write_bytes(b"not an xlsx workbook")
    portal = Mock()
    portal.ping.return_value = True
    portal.app = APP_FOURFRONT
    portal.server = "https://example.test"

    with patch("submitr.submission._define_portal", return_value=portal), \
            patch("submitr.submission._get_user_record", return_value={}), \
            patch("submitr.submission._resolve_app_args", return_value={
                "submission_centers": [], "consortia": [],
            }), \
            patch("submitr.submission._do_app_arg_defaulting", return_value=True), \
            patch("submitr.submission._get_submission_centers", return_value=[]), \
            patch("submitr.submission._is_admin_user", return_value=False), \
            patch("submitr.submission.get_metadata_bundles_bucket_from_health_path", return_value="bucket"):
        with pytest.raises(ProtectedDonorTransformError, match="Unable to load ProtectedDonor workbook") as error:
            submit_any_ingestion(
                ingestion_filename=str(source),
                ingestion_type="metadata_bundle",
                server="https://example.test",
                env="test",
                app=APP_FOURFRONT,
                noversion=True,
                no_query=True,
                **{mode: True},
            )

    assert error.value.__cause__ is not None
